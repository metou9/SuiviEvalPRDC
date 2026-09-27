from datetime import datetime
from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.utils import timezone
from django.utils.dateparse import parse_datetime

from openpyxl import load_workbook

from apps.geo.models import GeoUnit
from apps.infrastructure.models import Infrastructure


MAIN_SHEET = "Infrastructure - 3"


# ======================================================================
# OUTILS GENERAUX
# ======================================================================

def clean(value):
    """
    Nettoie une valeur provenant d'Excel.
    """
    if value is None:
        return ""

    if isinstance(value, str):
        return value.strip()

    return value


def to_decimal(value):
    if value in (None, ""):
        return None

    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError, TypeError):
        return None


def to_int(value):
    if value in (None, ""):
        return None

    try:
        return int(float(value))
    except (ValueError, TypeError):
        return None


def to_bool(value):
    """
    Conversion des réponses Kobo oui/non vers True/False.

    Retourne None si la valeur n'est pas renseignée.
    """
    if value in (None, ""):
        return None

    value = str(value).strip().lower()

    if value in (
        "oui",
        "yes",
        "true",
        "1",
    ):
        return True

    if value in (
        "non",
        "no",
        "false",
        "0",
    ):
        return False

    return None


def to_date(value):
    """
    Convertit une valeur Excel/Kobo en date.
    """
    if not value:
        return None

    if isinstance(value, datetime):
        return value.date()

    try:
        return datetime.fromisoformat(
            str(value)
        ).date()

    except (ValueError, TypeError):
        return None


def to_datetime(value):
    """
    Convertit la date/heure Kobo en datetime timezone-aware.

    Ceci évite les warnings Django concernant les dates naïves.
    """
    if not value:
        return None

    if isinstance(value, datetime):
        dt = value
    else:
        dt = parse_datetime(
            str(value)
        )

    if dt is None:
        return None

    if timezone.is_naive(dt):
        dt = timezone.make_aware(
            dt,
            timezone.get_current_timezone(),
        )

    return dt


def split_multiple(value):
    """
    Transforme un select_multiple Kobo en liste.

    Exemple :
    "Présidente Trésorière"
    devient :
    ["Présidente", "Trésorière"]
    """
    if not value:
        return []

    if isinstance(value, list):
        return value

    return [
        item.strip()
        for item in str(value)
        .replace(",", " ")
        .split()
        if item.strip()
    ]


# ======================================================================
# GPS
# ======================================================================

def parse_gps(row):
    """
    Priorité aux colonnes GPS séparées générées par Kobo.

    Sinon lecture de la colonne "Coordonnée GPS".
    """

    latitude = to_decimal(
        row.get(
            "_Coordonnée GPS_latitude"
        )
    )

    longitude = to_decimal(
        row.get(
            "_Coordonnée GPS_longitude"
        )
    )

    altitude = to_decimal(
        row.get(
            "_Coordonnée GPS_altitude"
        )
    )

    accuracy = to_decimal(
        row.get(
            "_Coordonnée GPS_precision"
        )
    )

    if (
        latitude is not None
        and longitude is not None
    ):
        return (
            latitude,
            longitude,
            altitude,
            accuracy,
        )

    gps = clean(
        row.get(
            "Coordonnée GPS"
        )
    )

    if not gps:
        return (
            None,
            None,
            None,
            None,
        )

    try:
        parts = str(gps).split()

        latitude = (
            to_decimal(parts[0])
            if len(parts) > 0
            else None
        )

        longitude = (
            to_decimal(parts[1])
            if len(parts) > 1
            else None
        )

        altitude = (
            to_decimal(parts[2])
            if len(parts) > 2
            else None
        )

        accuracy = (
            to_decimal(parts[3])
            if len(parts) > 3
            else None
        )

        return (
            latitude,
            longitude,
            altitude,
            accuracy,
        )

    except Exception:
        return (
            None,
            None,
            None,
            None,
        )


# ======================================================================
# REFERENTIEL GEOGRAPHIQUE
# ======================================================================

def find_geo_unit(
    project,
    commune,
):
    """
    Recherche la commune Kobo dans le référentiel géographique
    PRDC-VFS existant.

    IMPORTANT :
    - aucune commune n'est créée automatiquement ;
    - les valeurs Kobo originales restent conservées ;
    - les alias servent uniquement à retrouver une commune
      existante dans le référentiel.
    """

    if not commune:
        return None

    commune = str(
        commune
    ).strip()

    # --------------------------------------------------------------
    # Recherche exacte
    # --------------------------------------------------------------

    geo = (
        GeoUnit.objects
        .filter(
            project=project,
            geo_level__code="COM",
            name__iexact=commune,
        )
        .first()
    )

    if geo:
        return geo

    # --------------------------------------------------------------
    # Différences connues Kobo / référentiel
    # --------------------------------------------------------------

    aliases = {
        "lexeiba 2": "Lexeiba",
        "shamama": "Chemama",
    }

    mapped_name = aliases.get(
        commune.lower()
    )

    if mapped_name:
        return (
            GeoUnit.objects
            .filter(
                project=project,
                geo_level__code="COM",
                name__iexact=mapped_name,
            )
            .first()
        )

    # --------------------------------------------------------------
    # Ne jamais créer automatiquement une commune inconnue
    # --------------------------------------------------------------

    return None


# ======================================================================
# LECTURE EXCEL
# ======================================================================

def rows_as_dicts(sheet):
    """
    Retourne chaque ligne Excel sous forme de dictionnaire :

    {
        "Wilaya": "...",
        "Commune": "...",
        ...
    }
    """

    rows = sheet.iter_rows(
        values_only=True
    )

    try:
        headers = [
            clean(value)
            for value in next(rows)
        ]

    except StopIteration:
        return

    for values in rows:

        row = dict(
            zip(
                headers,
                values,
            )
        )

        if any(
            value not in (
                None,
                "",
            )
            for value in values
        ):
            yield row


# ======================================================================
# VOLETS
# ======================================================================

def get_volets(row):
    """
    Détermine les volets renseignés pour l'infrastructure.

    Le fichier Kobo contient :
    - la valeur globale du select_multiple ;
    - une colonne booléenne par choix.

    On utilise les deux pour rendre l'import robuste.
    """

    volets_value = str(
        clean(
            row.get(
                "Quels volets souhaitez-vous renseigner pour cette infrastructure ?"
            )
        )
    ).lower()

    # --------------------------------------------------------------
    # Réhabilitation
    # --------------------------------------------------------------

    rehabilitation_column = to_bool(
        row.get(
            "Quels volets souhaitez-vous renseigner pour cette infrastructure ?/Réhabilitation ou mise à niveau des infrastructures"
        )
    )

    has_rehabilitation = (
        rehabilitation_column is True
        or "réhabilitation" in volets_value
        or "rehabilitation" in volets_value
        or "mise à niveau" in volets_value
    )

    # --------------------------------------------------------------
    # Maintenance
    # --------------------------------------------------------------

    maintenance_column = to_bool(
        row.get(
            "Quels volets souhaitez-vous renseigner pour cette infrastructure ?/Fonctionnement et maintenance des infrastructures"
        )
    )

    has_maintenance = (
        maintenance_column is True
        or "maintenance" in volets_value
        or "fonctionnement" in volets_value
    )

    # --------------------------------------------------------------
    # Gouvernance
    # --------------------------------------------------------------

    governance_column = to_bool(
        row.get(
            "Quels volets souhaitez-vous renseigner pour cette infrastructure ?/Gouvernance et participation des femmes"
        )
    )

    has_governance = (
        governance_column is True
        or "gouvernance" in volets_value
        or "participation des femmes" in volets_value
    )

    return (
        has_rehabilitation,
        has_maintenance,
        has_governance,
    )


# ======================================================================
# FONCTIONS OCCUPEES PAR LES FEMMES
# ======================================================================

def get_women_functions(row):
    """
    Récupère les fonctions occupées par les femmes.

    Le fichier Kobo contient :
    - la valeur globale du select_multiple ;
    - les colonnes individuelles par fonction.
    """

    value = clean(
        row.get(
            "6. Fonctions occupées par les femmes dans le comité"
        )
    )

    functions = split_multiple(
        value
    )

    # --------------------------------------------------------------
    # Sécurité supplémentaire :
    # lecture des colonnes individuelles Kobo
    # --------------------------------------------------------------

    choices = [
        (
            "Présidente",
            "6. Fonctions occupées par les femmes dans le comité/Présidente",
        ),
        (
            "Vice-Présidente",
            "6. Fonctions occupées par les femmes dans le comité/Vice-Présidente",
        ),
        (
            "Trésorière",
            "6. Fonctions occupées par les femmes dans le comité/Trésorière",
        ),
        (
            "Secrétaire",
            "6. Fonctions occupées par les femmes dans le comité/Secrétaire",
        ),
        (
            "Membre simple",
            "6. Fonctions occupées par les femmes dans le comité/Membre simple",
        ),
        (
            "Autre",
            "6. Fonctions occupées par les femmes dans le comité/Autre",
        ),
    ]

    for label, column in choices:

        selected = to_bool(
            row.get(column)
        )

        if (
            selected is True
            and label not in functions
        ):
            functions.append(
                label
            )

    return functions


# ======================================================================
# IMPORT PRINCIPAL
# ======================================================================

@transaction.atomic
def import_kobo_infrastructure(
    file_obj,
    project,
    user=None,
):
    """
    Importe l'export Excel KoboToolbox Infrastructure.

    Une soumission est identifiée par :

        project + _uuid

    Un fichier peut donc être réimporté sans créer de doublons :
    les infrastructures existantes sont mises à jour.
    """

    workbook = load_workbook(
        file_obj,
        data_only=True,
    )

    # --------------------------------------------------------------
    # Vérification de la feuille
    # --------------------------------------------------------------

    if MAIN_SHEET not in workbook.sheetnames:
        raise ValueError(
            f"La feuille '{MAIN_SHEET}' "
            f"est introuvable dans le fichier Kobo."
        )

    sheet = workbook[
        MAIN_SHEET
    ]

    created_count = 0
    updated_count = 0
    ignored_count = 0

    geo_not_found = []


    # ==================================================================
    # LECTURE DES SOUMISSIONS
    # ==================================================================

    for row in rows_as_dicts(
        sheet
    ):

        # --------------------------------------------------------------
        # UUID KOBO
        # --------------------------------------------------------------

        kobo_uuid = clean(
            row.get(
                "_uuid"
            )
        )

        if not kobo_uuid:
            ignored_count += 1
            continue


        # --------------------------------------------------------------
        # GEOGRAPHIE
        # --------------------------------------------------------------

        wilaya = clean(
            row.get(
                "Wilaya"
            )
        )

        moughataa = clean(
            row.get(
                "Moughataa"
            )
        )

        commune = clean(
            row.get(
                "Commune"
            )
        )

        village_locality = clean(
            row.get(
                "Village/Localité :"
            )
        )

        geo_unit = find_geo_unit(
            project,
            commune,
        )


        # --------------------------------------------------------------
        # Commune non trouvée
        # --------------------------------------------------------------

        if (
            commune
            and not geo_unit
        ):
            geo_not_found.append(
                {
                    "uuid":
                        str(
                            kobo_uuid
                        ),

                    "wilaya":
                        str(
                            wilaya
                        ),

                    "moughataa":
                        str(
                            moughataa
                        ),

                    "commune":
                        str(
                            commune
                        ),
                }
            )


        # --------------------------------------------------------------
        # GPS
        # --------------------------------------------------------------

        (
            latitude,
            longitude,
            altitude,
            accuracy,
        ) = parse_gps(
            row
        )


        # --------------------------------------------------------------
        # VOLETS
        # --------------------------------------------------------------

        (
            has_rehabilitation,
            has_maintenance,
            has_governance,
        ) = get_volets(
            row
        )


        # --------------------------------------------------------------
        # FONCTIONS DES FEMMES
        # --------------------------------------------------------------

        women_functions = (
            get_women_functions(
                row
            )
        )


        # ==================================================================
        # MAPPING KOBO -> MODELE DJANGO
        # ==================================================================

        defaults = {

            # ==============================================================
            # TRACABILITE KOBO
            # ==============================================================

            "kobo_id":
                str(
                    clean(
                        row.get(
                            "_id"
                        )
                    )
                ),

            "kobo_submission_time":
                to_datetime(
                    row.get(
                        "_submission_time"
                    )
                ),

            "kobo_submitted_by":
                str(
                    clean(
                        row.get(
                            "_submitted_by"
                        )
                    )
                ),


            # ==============================================================
            # GEOGRAPHIE
            # ==============================================================

            "geo_unit":
                geo_unit,

            "kobo_wilaya":
                str(
                    wilaya
                ),

            "kobo_moughataa":
                str(
                    moughataa
                ),

            "kobo_commune":
                str(
                    commune
                ),

            "village_locality":
                str(
                    village_locality
                ),


            # ==============================================================
            # GPS
            # ==============================================================

            "gps_latitude":
                latitude,

            "gps_longitude":
                longitude,

            "gps_altitude":
                altitude,

            "gps_accuracy":
                accuracy,


            # ==============================================================
            # INFORMATIONS GENERALES
            # ==============================================================

            "enumerator_name":
                str(
                    clean(
                        row.get(
                            "Nom de l’enquêteur"
                        )
                    )
                ),

            "name":
                str(
                    clean(
                        row.get(
                            "Nom / désignation de l'infrastructure"
                        )
                    )
                ),

            "infrastructure_type":
                str(
                    clean(
                        row.get(
                            "Type d’infrastructure concernée"
                        )
                    )
                ),

            "infrastructure_type_other":
                str(
                    clean(
                        row.get(
                            "Précisez le type d’infrastructure"
                        )
                    )
                ),


            # ==============================================================
            # VOLETS
            # ==============================================================

            "has_rehabilitation":
                has_rehabilitation,

            "has_maintenance":
                has_maintenance,

            "has_governance":
                has_governance,


            # ==============================================================
            # REHABILITATION / MISE A NIVEAU
            # ==============================================================

            "intervention_type":
                str(
                    clean(
                        row.get(
                            "1. Nature de l’intervention réalisée"
                        )
                    )
                ),

            "intervention_type_other":
                str(
                    clean(
                        row.get(
                            "Précisez la nature de l’intervention"
                        )
                    )
                ),

            "works_completed":
                to_bool(
                    row.get(
                        "2. Les travaux sont-ils achevés ?"
                    )
                ),

            "completion_date":
                to_date(
                    row.get(
                        "3. Date d’achèvement des travaux"
                    )
                ),

            "rehab_functional":
                to_bool(
                    row.get(
                        "4. L’infrastructure est-elle actuellement fonctionnelle ?"
                    )
                ),


            # ==============================================================
            # FONCTIONNEMENT / MAINTENANCE
            # ==============================================================

            "management_structure_exists":
                to_bool(
                    row.get(
                        "1. Une structure de gestion de l’infrastructure est-elle mise en place"
                    )
                ),

            "management_structure_name":
                str(
                    clean(
                        row.get(
                            "2. Nom de la structure"
                        )
                    )
                ),

            "maintenance_functional":
                to_bool(
                    row.get(
                        "3. L’infrastructure est-elle actuellement fonctionnelle ?"
                    )
                ),

            "maintenance_last_12_months":
                to_bool(
                    row.get(
                        "4. Des opérations de maintenance ont-elles été réalisées au cours des 12 derniers mois ?"
                    )
                ),


            # ==============================================================
            # GOUVERNANCE
            # ==============================================================

            "committee_name":
                str(
                    clean(
                        row.get(
                            "1. Nom de l’organe ou comité concerné"
                        )
                    )
                ),

            "committee_functional":
                to_bool(
                    row.get(
                        "2. Le comité est-il actuellement fonctionnel ?"
                    )
                ),

            "committee_has_women":
                to_bool(
                    row.get(
                        "3. Le comité comprend-il des femmes ?"
                    )
                ),

            "committee_members_total":
                to_int(
                    row.get(
                        "4. Nombre total de membres du comité"
                    )
                ),

            "committee_women_total":
                to_int(
                    row.get(
                        "5. Nombre de femmes membres"
                    )
                ),

            "women_functions":
                women_functions,

            "women_functions_other":
                str(
                    clean(
                        row.get(
                            "Précisez les fonctions occupées"
                        )
                    )
                ),
        }


        # --------------------------------------------------------------
        # UTILISATEUR AYANT EFFECTUE L'IMPORT
        # --------------------------------------------------------------

        if user:
            defaults[
                "updated_by"
            ] = user


        # ==============================================================
        # CREATION / MISE A JOUR
        # ==============================================================

        (
            infrastructure,
            created,
        ) = (
            Infrastructure.objects
            .update_or_create(
                project=project,
                kobo_uuid=str(
                    kobo_uuid
                ),
                defaults=defaults,
            )
        )


        # --------------------------------------------------------------
        # COMPTEURS
        # --------------------------------------------------------------

        if created:
            created_count += 1

            if user:
                infrastructure.created_by = (
                    user
                )

                infrastructure.save(
                    update_fields=[
                        "created_by"
                    ]
                )

        else:
            updated_count += 1


    # ==================================================================
    # RESULTAT
    # ==================================================================

    return {
        "created":
            created_count,

        "updated":
            updated_count,

        "ignored":
            ignored_count,

        "geo_not_found":
            geo_not_found,

        "geo_not_found_count":
            len(
                geo_not_found
            ),
    }